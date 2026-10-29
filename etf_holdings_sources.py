"""Fetch dated ETF constituent snapshots from public issuer data sources.

Common European funds use the issuers' public APIs directly so the app does
not depend on an optional third-party scraper being installed. The optional
adapters remain as a fallback for other supported issuers.
"""
from __future__ import annotations

from datetime import date, datetime, timezone
from functools import lru_cache
import csv
from io import BytesIO, StringIO
import re
from typing import Any

import pandas as pd
import requests


# Known portfolio listings are resolved locally before asking Yahoo for an ISIN.
# These identify share classes, not holdings inside a fund.
KNOWN_FUND_ISINS = {
    "EXUS.DE": "IE0006WW1TQ4",
    "VWCE.DE": "IE00BK5BQT80",
    "2B70.DE": "IE00BYXG2H39",
    "36BZ.DE": "IE00BQT3WG13",
    "EUNK.DE": "IE00B4K48X80",
    "IB1T.DE": "XS2940466316",
    "INDA": "US46429B5984",
    "IS3N.DE": "IE00BKM4GZ66",
    "IUSQ.DE": "IE00B6R52259",
    "LYP6.DE": "LU0908500753",
    "MCHI": "US46429B6719",
    "PPFB.DE": "IE00B4ND3602",
    "QDV5.DE": "IE00BZCQB185",
    "QDVG.DE": "IE00B43HR379",
    "SPYL.DE": "IE000XZSV718",
    "SXR8.DE": "IE00B5BMR087",
}

# Route mapped share classes straight to their issuer instead of trying
# unrelated provider catalogues (which can send valid funds to the wrong API).
KNOWN_ISIN_ISSUERS = {
    "IE0006WW1TQ4": "Xtrackers",
    "IE00BK5BQT80": "Vanguard",
    "IE00BYXG2H39": "iShares",
    "IE00BQT3WG13": "iShares",
    "IE00B4K48X80": "iShares",
    "IE00BKM4GZ66": "iShares",
    "IE00B6R52259": "iShares",
    "LU0908500753": "Amundi",
    "IE00BZCQB185": "iShares",
    "IE00B43HR379": "iShares",
    "IE00B5BMR087": "iShares",
}

# iShares US funds and State Street's SPDR use official issuer CSV/XLSX
# downloads rather than the Italy-specific adapters in portfolio-scraper.
ISHARES_US_HOLDINGS = {
    "INDA": "https://www.ishares.com/us/products/239659/ishares-msci-india-etf/latest-holdings.csv",
    "MCHI": "https://www.ishares.com/us/products/239619/ishares-msci-china-etf/latest-holdings.csv",
}
SSGA_SPYL_HOLDINGS_URL = "https://www.ssga.com/library-content/products/fund-data/etfs/emea/holdings-daily-emea-en-spyl-gy.xlsx"
XTRACKERS_HOLDINGS_URL = "https://etf.dws.com/etfdata/export/ITA/ITA/csv/product/constituent/{isin}/"
VANGUARD_EU_GRAPHQL_URL = "https://www.ie.vanguard/gpx/graphql"
VANGUARD_EU_CONSUMER_ID = "ie0"
VANGUARD_EU_PORT_IDS = {
    "IE00BK5BQT80": "9679",  # FTSE All-World UCITS ETF (USD) Accumulating
}
VANGUARD_EU_PRODUCT_URLS = {
    "IE00BK5BQT80": "https://www.nl.vanguard/professional/product/etf/equity/9679/ftse-all-world-ucits",
}
VANGUARD_EU_EQUITY_TYPES = [
    "EQ.DRCPT", "EQ.ETF", "EQ.FSH", "EQ.PREF", "EQ.PSH", "EQ.REIT",
    "EQ.RIGHT", "EQ.STOCK", "EQ.WRT",
]
VANGUARD_EU_HOLDINGS_QUERY = """
query FundsHoldingsQuery($portIds: [String!], $securityTypes: [String!], $lastItemKey: String) {
  funds(portIds: $portIds) {
    profile { fundFullName fundCurrency primarySectorEquityClassification }
  }
  borHoldings(portIds: $portIds) {
    holdings(limit: 1500, securityTypes: $securityTypes, lastItemKey: $lastItemKey) {
      items {
        issuerName securityLongDescription gicsSectorDescription icbSectorDescription
        icbIndustryDescription marketValuePercentage sedol1 quantity ticker securityType
        effectiveDate marketValueBaseCurrency bloombergIsoCountry
      }
      totalHoldings lastItemKey
    }
  }
}
"""

# These are exchange-traded products, not share-holding ETFs. Their issuer
# documentation identifies one underlying asset, so show that explicitly in
# X-Ray instead of trying to download company constituents.
SINGLE_ASSET_EXPOSURES = {
    "IB1T.DE": ("Bitcoin", "BTC", "Digital assets", "https://www.ishares.com/de/privatanleger/de/produkte/337088/ishares-bitcoin-etp"),
    "PPFB.DE": ("Physical gold", "GOLD", "Commodities", "https://www.ishares.com/de/privatanleger/de/produkte/258441/"),
}

ISSUER_HOME = {
    "Amundi": "https://www.amundietf.com/",
    "iShares": "https://www.ishares.com/",
    "Vanguard": "https://global.vanguard.com/",
    "Xtrackers": "https://etf.dws.com/",
}


def _resolved_isin(ticker: str, supplied_isin: str = "") -> str:
    isin = str(supplied_isin or "").strip().upper()
    if _valid_isin(isin):
        return isin
    ticker = str(ticker or "").strip().upper()
    if ticker in KNOWN_FUND_ISINS:
        return KNOWN_FUND_ISINS[ticker]
    try:
        import yfinance as yf

        instrument = yf.Ticker(ticker)
        isin = str(getattr(instrument, "isin", "") or "").strip().upper()
        if _valid_isin(isin):
            return isin
        info = instrument.info or {}
        isin = str(info.get("isin", "") or "").strip().upper()
        return isin if _valid_isin(isin) else ""
    except Exception:
        return ""


def _valid_isin(value: str) -> bool:
    return bool(re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}[0-9]", str(value or "").strip().upper()))


def _parse_weight_fraction(value: Any, percentage_points: bool = False) -> float:
    """Convert an issuer weight to a fraction, accepting either 0.08 or 8.0% formats."""
    if value is None or value is pd.NA:
        return float("nan")
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    # European issuer files commonly use comma decimals. A comma-only value
    # cannot be a grouped thousands value for a portfolio percentage, so read
    # it as a decimal separator. US-style grouping commas are removed when a
    # decimal point is also present.
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    else:
        text = text.replace(",", "")
    has_percent = text.endswith("%")
    if has_percent:
        text = text[:-1].strip()
    try:
        number = float(text)
    except (TypeError, ValueError):
        return float("nan")
    if has_percent or percentage_points:
        number /= 100.0
    return number


def _direct_snapshot(
    fund: str,
    isin: str,
    issuer: str,
    frame: pd.DataFrame,
    *,
    source_url: str,
    as_of: str = "",
    equity_only: bool = False,
) -> dict[str, Any]:
    """Convert a direct issuer download into the app's saved snapshot format."""
    name_col = _pick_column(frame, (
        "name", "holding_name", "security_name", "issuer_name", "security_long_description",
        "company_name", "nome", "denominazione", "titolo",
    ))
    ticker_col = _pick_column(frame, ("ticker", "symbol", "sedol", "sedol1", "isin", "codice"))
    weight_col = _pick_column(frame, (
        "weight", "weight_pct", "weight_percent", "weight_%", "percent_of_fund",
        "market_value_percentage", "percent_of_market_value", "portfolio_weight",
        "weight_in_percent", "weight_in_%", "portfolio_weighting", "weighting_pct",
        "peso", "percentuale",
    ))
    type_col = _pick_column(frame, ("type", "asset_class", "security_type"))
    if name_col is None or weight_col is None:
        raise ValueError("the issuer download did not contain recognizable name and weight columns")
    if equity_only and type_col:
        type_values = frame[type_col].fillna("").astype(str).str.strip().str.upper()
        equity_rows = type_values.isin({"EQUITY", "STOCK"}) | type_values.str.startswith("EQ.")
        if equity_rows.any():
            frame = frame.loc[equity_rows].copy()

    tickers = frame[ticker_col] if ticker_col else pd.Series("", index=frame.index)
    rows = pd.DataFrame(index=frame.index)
    rows["constituent_name"] = frame[name_col].fillna("").astype(str).str.strip()
    rows["constituent_ticker"] = tickers.fillna("").astype(str).str.strip().str.upper()
    no_ticker = rows["constituent_ticker"].eq("")
    rows.loc[no_ticker, "constituent_ticker"] = rows.loc[no_ticker, "constituent_name"].str.upper()
    # Issuer CSVs supply percentage points (e.g. 6.21); the shared path below
    # expects a fraction (0.0621) and converts it to percentage points once.
    weight_values = frame[weight_col].map(lambda value: _parse_weight_fraction(value))
    weight_labels = frame[weight_col].fillna("").astype(str)
    numeric_weights = pd.to_numeric(weight_values, errors="coerce").dropna()
    raw_total = float(numeric_weights.clip(lower=0).sum()) if not numeric_weights.empty else 0.0
    percentage_points = bool(weight_labels.str.contains("%", regex=False).any()) or (
        90.0 <= raw_total <= 110.0
        or (not numeric_weights.empty and raw_total > 1.0 and numeric_weights.abs().max() > 1.0)
    )
    rows["weight"] = frame[weight_col].map(lambda value: _parse_weight_fraction(value, percentage_points))
    sector_col = _pick_column(frame, (
        "sector", "sector_classification", "gics_sector_description", "icb_sector_description",
        "icb_industry_description", "settore",
    ))
    country_col = _pick_column(frame, (
        "country", "country_name", "country_code", "country_of_risk", "country_of_domicile",
        "trade_country_name", "location", "bloomberg_iso_country", "paese", "paese_di_rischio",
    ))
    region_col = _pick_column(frame, ("region",))
    rows["sector"] = frame[sector_col].map(_source_text) if sector_col else ""
    rows["country"] = frame[country_col].map(_source_text) if country_col else ""
    rows["region"] = frame[region_col].map(_source_text) if region_col else ""
    rows = rows.dropna(subset=["weight"])
    rows = rows[(rows["constituent_ticker"] != "") & rows["weight"].ge(0)]
    if rows.empty:
        raise ValueError("the issuer download contained no usable holdings rows")
    rows = rows.groupby("constituent_ticker", as_index=False, sort=False).agg(
        constituent_name=("constituent_name", lambda values: next((v for v in values if v), "")),
        weight=("weight", "sum"),
        sector=("sector", lambda values: next((v for v in values if v), "")),
        country=("country", lambda values: next((v for v in values if v), "")),
        region=("region", lambda values: next((v for v in values if v), "")),
    )
    rows["weight_pct"] = rows.pop("weight") * 100.0
    source_date = as_of or date.today().isoformat()
    fetched_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    return {
        "fund_ticker": fund,
        "fund_isin": isin,
        "holdings": rows,
        "source_name": f"{issuer} issuer holdings · ISIN {isin}",
        "source_url": source_url,
        "as_of": source_date,
        "fetched_at": fetched_at,
        "issuer": issuer,
        "holding_count": int(len(rows)),
    }


def _fetch_ishares_us_holdings(fund: str, isin: str) -> dict[str, Any]:
    url = ISHARES_US_HOLDINGS[fund]
    response = requests.get(url, timeout=30, headers={"User-Agent": "InvestmentCopilot/1.0 (public issuer holdings)"})
    response.raise_for_status()
    text = response.content.decode("utf-8-sig", errors="replace")
    lines = text.splitlines()
    header_index = next((
        index for index, line in enumerate(lines[:30])
        if {"ticker", "name"}.issubset({re.sub(r"[^a-z0-9]", "", cell.strip().strip('"').lower()) for cell in next(csv.reader([line]))})
        and any("weight" in re.sub(r"[^a-z0-9]", "", cell.strip().strip('"').lower()) for cell in next(csv.reader([line])))
    ), None)
    if header_index is None:
        raise ValueError("iShares CSV did not contain a holdings header")
    frame = pd.read_csv(StringIO("\n".join(lines[header_index:])))
    as_of = ""
    if len(lines) > 1:
        try:
            first_row = next(csv.reader([lines[1]]))
            parsed_date = pd.to_datetime(first_row[-1], errors="coerce")
            if not pd.isna(parsed_date):
                as_of = parsed_date.date().isoformat()
        except Exception:
            pass
    return _direct_snapshot(fund, isin, "iShares", frame, source_url=url, as_of=as_of, equity_only=True)


def _fetch_ssga_spyl(fund: str, isin: str) -> dict[str, Any]:
    response = requests.get(SSGA_SPYL_HOLDINGS_URL, timeout=30, headers={"User-Agent": "InvestmentCopilot/1.0 (public issuer holdings)"})
    response.raise_for_status()
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read State Street's official holdings workbook") from exc
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    data = list(sheet.iter_rows(values_only=True))
    header_index = None
    for index, cells in enumerate(data[:40]):
        normalized = {re.sub(r"[^a-z0-9]", "", str(cell or "").lower()) for cell in cells}
        has_weight = any("weight" in value or value in {"percentoffund", "percentofnetassets"} for value in normalized)
        has_security_name = any(value == "name" or "securityname" in value or value == "holdingname" for value in normalized)
        if has_weight and has_security_name:
            header_index = index
            break
    if header_index is None:
        raise ValueError("State Street workbook did not contain a recognizable holdings header")
    headers = [str(value or "").strip() for value in data[header_index]]
    frame = pd.DataFrame(data[header_index + 1:], columns=headers).dropna(how="all")
    return _direct_snapshot(fund, isin, "State Street", frame, source_url=SSGA_SPYL_HOLDINGS_URL, equity_only=True)


def _csv_holdings_frame(text: str, *, max_header_lines: int = 60) -> pd.DataFrame:
    """Find and parse a CSV header even when an issuer adds a title preamble."""
    lines = text.lstrip("\ufeff").splitlines()
    name_headers = {
        "name", "holdingname", "securityname", "issuername", "securitylongdescription",
        "companyname", "nome", "denominazione", "titolo",
    }
    weight_headers = {
        "weight", "weightpct", "weightpercent", "weightinpercent", "portfolioweight",
        "marketvaluepercentage", "percentofmarketvalue", "percentofmarketvalue",
        "peso", "percentuale",
    }
    for line_index, line in enumerate(lines[:max_header_lines]):
        delimiter = max((",", ";", "\t", "|"), key=line.count)
        if line.count(delimiter) < 1:
            continue
        try:
            headers = next(csv.reader([line], delimiter=delimiter))
        except (csv.Error, StopIteration):
            continue
        normalized = {
            re.sub(r"[^a-z0-9]", "", str(value).strip().strip('"').lower())
            for value in headers
        }
        has_name = bool(normalized & name_headers) or any("holdingname" in value for value in normalized)
        has_weight = bool(normalized & weight_headers) or any(
            "weight" in value or "percentofmarketvalue" in value for value in normalized
        )
        if has_name and has_weight:
            frame = pd.read_csv(StringIO("\n".join(lines[line_index:])), sep=delimiter, dtype=str)
            frame.columns = [str(column).strip() for column in frame.columns]
            return frame.dropna(how="all")
    raise ValueError("issuer CSV did not contain recognizable name and weight columns")


def _extract_source_date(text: str) -> str:
    """Read an explicit as-of date from an issuer file preamble, if present."""
    for line in text.splitlines()[:20]:
        if not re.search(r"as\s*of|holdings|portfolio|data|date|aggiorn|rifer", line, re.I):
            continue
        match = re.search(r"\b(20\d{2}[-/.]\d{1,2}[-/.]\d{1,2}|\d{1,2}[-/.]\d{1,2}[-/.]20\d{2})\b", line)
        if match:
            parsed = pd.to_datetime(match.group(1), errors="coerce", dayfirst=True)
            if not pd.isna(parsed):
                return parsed.date().isoformat()
    return ""


def _fetch_xtrackers_holdings(fund: str, isin: str) -> dict[str, Any]:
    """Download the official DWS/Xtrackers constituent CSV by ISIN."""
    url = XTRACKERS_HOLDINGS_URL.format(isin=isin)
    response = requests.get(
        url,
        timeout=30,
        headers={"User-Agent": "InvestmentCopilot/1.0 (public issuer holdings)"},
    )
    response.raise_for_status()
    text = response.content.decode("utf-8-sig", errors="replace")
    frame = _csv_holdings_frame(text)
    as_of = _source_as_of(frame)
    if as_of == date.today().isoformat():
        as_of = _extract_source_date(text) or as_of
    return _direct_snapshot(fund, isin, "Xtrackers", frame, source_url=url, as_of=as_of, equity_only=True)


def _fetch_vanguard_eu_holdings(fund: str, isin: str) -> dict[str, Any]:
    """Download all published Vanguard EU equity rows through its public GraphQL API."""
    port_id = VANGUARD_EU_PORT_IDS.get(isin)
    if not port_id:
        raise ValueError(f"No Vanguard EU product id is mapped for ISIN {isin}")

    headers = {
        "Content-Type": "application/json",
        "X-Consumer-ID": VANGUARD_EU_CONSUMER_ID,
        "User-Agent": "InvestmentCopilot/1.0 (public issuer holdings)",
    }
    rows: list[dict[str, Any]] = []
    profile: dict[str, Any] = {}
    last_item_key: str | None = None
    pagination_complete = False
    for _ in range(50):
        response = requests.post(
            VANGUARD_EU_GRAPHQL_URL,
            json={
                "query": VANGUARD_EU_HOLDINGS_QUERY,
                "variables": {
                    "portIds": [port_id],
                    "securityTypes": VANGUARD_EU_EQUITY_TYPES,
                    "lastItemKey": last_item_key,
                },
            },
            timeout=30,
            headers=headers,
        )
        response.raise_for_status()
        payload = response.json()
        if payload.get("errors"):
            messages = "; ".join(str(item.get("message", "GraphQL error")) for item in payload["errors"])
            raise ValueError(f"Vanguard holdings API: {messages[:500]}")
        data = payload.get("data") or {}
        funds = data.get("funds") or []
        if funds and isinstance(funds[0], dict):
            profile = profile or (funds[0].get("profile") or {})
        provider_holdings = data.get("borHoldings") or []
        if not provider_holdings or not isinstance(provider_holdings[0], dict):
            break
        holdings = provider_holdings[0].get("holdings") or {}
        page_rows = holdings.get("items") or []
        rows.extend(row for row in page_rows if isinstance(row, dict))
        next_key = holdings.get("lastItemKey")
        if not next_key or str(next_key) == str(last_item_key):
            pagination_complete = True
            break
        last_item_key = str(next_key)

    if not pagination_complete:
        raise ValueError("Vanguard holdings pagination did not finish; refusing to save an incomplete snapshot")
    if not rows:
        raise ValueError(f"Vanguard returned no equity holdings for product {port_id}")

    frame = pd.DataFrame(rows)
    sector_preference = str(profile.get("primarySectorEquityClassification") or "").strip().casefold()
    if sector_preference == "icb sectors":
        frame["sector"] = frame.get("icbSectorDescription", pd.Series("", index=frame.index)).fillna("")
        if "icbIndustryDescription" in frame:
            frame["sector"] = frame["sector"].where(frame["sector"].astype(str).str.strip() != "", frame["icbIndustryDescription"])
    else:
        frame["sector"] = frame.get("gicsSectorDescription", pd.Series("", index=frame.index))
    frame["name"] = frame.get("issuerName", pd.Series("", index=frame.index)).fillna("")
    if "securityLongDescription" in frame:
        frame["name"] = frame["name"].where(frame["name"].astype(str).str.strip() != "", frame["securityLongDescription"])
    ticker_values = frame.get("ticker", pd.Series("", index=frame.index)).fillna("").astype(str).str.strip()
    sedol_values = frame.get("sedol1", pd.Series("", index=frame.index)).fillna("").astype(str).str.strip()
    frame["ticker"] = ticker_values.where(ticker_values != "", sedol_values)
    frame["country"] = frame.get("bloombergIsoCountry", pd.Series("", index=frame.index))
    frame["weight"] = frame.get("marketValuePercentage", pd.Series(pd.NA, index=frame.index))
    source_url = VANGUARD_EU_PRODUCT_URLS.get(isin, ISSUER_HOME["Vanguard"])
    as_of = _source_as_of(frame)
    return _direct_snapshot(
        fund, isin, "Vanguard", frame, source_url=source_url, as_of=as_of, equity_only=True
    )


def _single_asset_snapshot(fund: str, isin: str) -> dict[str, Any]:
    name, symbol, category, source_url = SINGLE_ASSET_EXPOSURES[fund]
    rows = pd.DataFrame([{
        "constituent_ticker": symbol,
        "constituent_name": name,
        "weight_pct": 100.0,
        "sector": category,
        "country": "",
        "region": "",
    }])
    return {
        "fund_ticker": fund,
        "fund_isin": isin,
        "holdings": rows,
        "source_name": f"Issuer exposure · {name} · ISIN {isin}",
        "source_url": source_url,
        "as_of": date.today().isoformat(),
        "fetched_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "issuer": "iShares",
        "holding_count": 1,
    }


def _source_classes() -> list[tuple[str, Any]]:
    try:
        from portfolio_scraper.etf import (
            AmundiScraper,
            ISharesItScraper,
            VanguardItScraper,
            XTrackersItScraper,
        )
    except ImportError as exc:
        raise RuntimeError(
            "No optional fallback adapter is installed for this issuer. Built-in downloads cover Vanguard EU, Xtrackers, iShares US, and State Street SPYL."
        ) from exc
    return [
        ("Amundi", AmundiScraper),
        ("iShares", ISharesItScraper),
        ("Vanguard", VanguardItScraper),
        ("Xtrackers", XTrackersItScraper),
    ]


@lru_cache(maxsize=4)
def _issuer_listings(issuer: str, scraper_class: Any) -> pd.DataFrame:
    """Reuse issuer product catalogs while separately refreshing fund holdings."""
    listings = scraper_class().get_listings()
    return listings.copy() if isinstance(listings, pd.DataFrame) else pd.DataFrame()


def _pick_column(frame: pd.DataFrame, names: tuple[str, ...]) -> str | None:
    by_normalized = {re.sub(r"[^a-z0-9]", "", str(c).lower()): str(c) for c in frame.columns}
    for name in names:
        match = by_normalized.get(re.sub(r"[^a-z0-9]", "", name.lower()))
        if match is not None:
            return match
    return None


def _source_as_of(frame: pd.DataFrame) -> str:
    date_column = _pick_column(
        frame,
        ("as_of", "asof", "as_of_date", "holdings_date", "portfolio_date", "report_date", "effective_date"),
    )
    if date_column:
        dates = pd.to_datetime(frame[date_column], errors="coerce", utc=True).dropna()
        if not dates.empty:
            return dates.max().date().isoformat()
    return date.today().isoformat()


def _source_text(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    # The adapter's normalized sector fields can be Enum values.
    return str(getattr(value, "value", value)).strip()


def fetch_issuer_holdings(fund_ticker: str, isin: str = "") -> dict[str, Any]:
    """Fetch one ETF snapshot; return normalized data plus source/date metadata."""
    fund = str(fund_ticker or "").strip().upper()
    resolved_isin = _resolved_isin(fund, isin)
    if not resolved_isin:
        raise ValueError(
            f"Could not identify {fund} by ISIN. Add its ISIN in ETF Research → Instrument Master."
        )

    if fund in SINGLE_ASSET_EXPOSURES:
        return _single_asset_snapshot(fund, resolved_isin)
    if fund in ISHARES_US_HOLDINGS:
        return _fetch_ishares_us_holdings(fund, resolved_isin)
    if fund == "SPYL.DE":
        return _fetch_ssga_spyl(fund, resolved_isin)

    preferred_issuer = KNOWN_ISIN_ISSUERS.get(resolved_isin)
    direct_source_error = ""
    if preferred_issuer == "Xtrackers":
        try:
            return _fetch_xtrackers_holdings(fund, resolved_isin)
        except Exception as exc:
            direct_source_error = f"Xtrackers official download: {exc}"
    elif preferred_issuer == "Vanguard" and resolved_isin in VANGUARD_EU_PORT_IDS:
        try:
            return _fetch_vanguard_eu_holdings(fund, resolved_isin)
        except Exception as exc:
            direct_source_error = f"Vanguard official API: {exc}"

    try:
        all_sources = _source_classes()
    except RuntimeError as exc:
        if direct_source_error:
            raise ValueError(f"Could not load {fund} from {direct_source_error}") from exc
        raise
    # For a known ISIN use the matching provider only. For other funds, search
    # the supported providers in turn, using their catalog as the identifier
    # lookup.
    sources = (
        [source for source in all_sources if source[0] == preferred_issuer]
        if preferred_issuer
        else all_sources
    )
    failures = [direct_source_error] if direct_source_error else []
    for issuer, scraper_class in sources:
        try:
            scraper = scraper_class()
            listings = pd.DataFrame()
            listing = pd.Series(dtype=object)
            # Xtrackers' documented holdings identifier is the ISIN itself;
            # do not make an unnecessary catalog request for that provider.
            if issuer == "Xtrackers":
                internal_id = resolved_isin
            else:
                listings = _issuer_listings(issuer, scraper_class)
                isin_column = _pick_column(listings, ("isin",)) if not listings.empty else None
                if isin_column is None:
                    if preferred_issuer:
                        raise ValueError("the issuer product list did not include an ISIN column")
                    continue
                matched = listings[
                    listings[isin_column].fillna("").astype(str).str.strip().str.upper() == resolved_isin
                ]
                if matched.empty:
                    if preferred_issuer:
                        raise ValueError("the issuer product list did not include this fund ISIN")
                    continue
                listing = matched.iloc[0]
                internal_id = listing.get("internal_id", resolved_isin)
                if pd.isna(internal_id) or not str(internal_id).strip():
                    internal_id = resolved_isin

            # Use the scraper's normalized holdings response. Fetching both its
            # raw and normalized forms made the same issuer file download
            # twice, slowing first page load and increasing timeout risk.
            normalized = scraper.get_holdings(internal_id)
            if normalized is None or normalized.empty:
                raise ValueError("the issuer returned no holdings rows")

            name_col = _pick_column(normalized, ("name", "holding_name", "security_name"))
            ticker_col = _pick_column(normalized, ("ticker", "symbol"))
            isin_col = _pick_column(normalized, ("isin",))
            weight_col = _pick_column(normalized, ("weight", "weight_pct", "holding_percent"))
            sector_col = _pick_column(normalized, ("sector",))
            country_col = _pick_column(normalized, ("country",))
            region_col = _pick_column(normalized, ("region",))
            if weight_col is None:
                raise ValueError("the issuer response did not include holding weights")

            rows = pd.DataFrame(index=normalized.index)
            names = normalized[name_col] if name_col else pd.Series("", index=normalized.index)
            tickers = normalized[ticker_col] if ticker_col else pd.Series("", index=normalized.index)
            holding_isins = normalized[isin_col] if isin_col else pd.Series("", index=normalized.index)
            rows["constituent_name"] = names.fillna("").astype(str).str.strip()
            rows["constituent_ticker"] = tickers.fillna("").astype(str).str.strip().str.upper()
            missing_id = rows["constituent_ticker"].eq("")
            rows.loc[missing_id, "constituent_ticker"] = holding_isins[missing_id].fillna("").astype(str).str.strip().str.upper()
            missing_id = rows["constituent_ticker"].eq("")
            rows.loc[missing_id, "constituent_ticker"] = rows.loc[missing_id, "constituent_name"].str.upper()
            rows["weight_pct"] = pd.to_numeric(normalized[weight_col], errors="coerce") * 100.0
            rows["sector"] = normalized[sector_col].map(_source_text) if sector_col else ""
            rows["country"] = normalized[country_col].map(_source_text) if country_col else ""
            rows["region"] = normalized[region_col].map(_source_text) if region_col else ""
            rows = rows.dropna(subset=["weight_pct"])
            rows = rows[(rows["constituent_ticker"] != "") & (rows["weight_pct"] >= 0)]
            if rows.empty:
                raise ValueError("the issuer response contained no usable constituent rows")

            # Providers can report multiple share lines for one company. Keep
            # their total weight while retaining the first available labels.
            rows = rows.groupby("constituent_ticker", as_index=False, sort=False).agg(
                constituent_name=("constituent_name", lambda values: next((v for v in values if v), "")),
                weight_pct=("weight_pct", "sum"),
                sector=("sector", lambda values: next((v for v in values if v), "")),
                country=("country", lambda values: next((v for v in values if v), "")),
                region=("region", lambda values: next((v for v in values if v), "")),
            )
            source_date = _source_as_of(normalized)
            fetched_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
            rows["_source_date"] = source_date
            url_column = _pick_column(listings, ("product_url", "fund_url", "url", "website")) if not listings.empty else None
            url_value = listing.get(url_column, "") if url_column else ""
            source_url = "" if pd.isna(url_value) else str(url_value).strip()
            if not source_url.startswith(("https://", "http://")):
                source_url = ISSUER_HOME.get(issuer, "")
            return {
                "fund_ticker": fund,
                "fund_isin": resolved_isin,
                "holdings": rows,
                "source_name": f"{issuer} issuer holdings · ISIN {resolved_isin}",
                "source_url": source_url,
                "as_of": source_date,
                "fetched_at": fetched_at,
                "issuer": issuer,
                "holding_count": int(len(rows)),
            }
        except Exception as exc:
            failures.append(f"{issuer}: {exc}")

    detail = "; ".join(failures[-3:])
    raise ValueError(
        f"Could not load issuer holdings for {fund} (ISIN {resolved_isin})."
        + (f" Source details: {detail}" if detail else "")
    )
